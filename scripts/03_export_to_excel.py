#!/usr/bin/env python3
"""
Use 3: Importing the API into Excel.

Downloads the public SafeSpace API and builds an Excel workbook:
  - "Quests" sheet: the raw API data (Quest, Category, User Count)
  - "By Category" sheet: totals per category, calculated with Excel
    SUMIF formulas, plus a bar chart
"""

import os
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill

API_URL = os.environ.get(
    "SAFESPACE_API_URL",
    "https://ilona.pythonanywhere.com/api/quest-summary/",
)
OUTPUT_PATH = Path(os.environ.get("SAFESPACE_OUTPUT_FILE", "safespace_quests.xlsx"))

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="324841")  # SafeSpace dark slate


def style_header(sheet):
    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL


def main():
    print(f"API: {API_URL}")

    try:
        response = requests.get(API_URL, timeout=10)
        response.raise_for_status()
    except requests.RequestException as exc:
        print(f"Could not reach the API: {exc}")
        return 1

    quests = response.json()
    print(f"Fetched {len(quests)} quests")

    workbook = Workbook()

    # Sheet 1: the API data as rows
    sheet = workbook.active
    sheet.title = "Quests"
    sheet.append(["Quest", "Category", "User Count"])
    for item in quests:
        sheet.append([item["quest"], item["category"], item["user_count"]])
    style_header(sheet)
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 22
    sheet.column_dimensions["C"].width = 12

    last_row = len(quests) + 1
    sheet.append([])
    sheet.append(["Total", "", f"=SUM(C2:C{last_row})"])
    sheet.append(["Average", "", f"=ROUND(AVERAGE(C2:C{last_row}), 1)"])

    # Sheet 2: totals per category, with Excel formulas and a chart
    summary = workbook.create_sheet("By Category")
    summary.append(["Category", "Total Users"])
    categories = sorted({item["category"] for item in quests})
    for row, category in enumerate(categories, start=2):
        summary.append([category, f"=SUMIF(Quests!B:B, A{row}, Quests!C:C)"])
    style_header(summary)
    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 14

    chart = BarChart()
    chart.title = "SafeSpace completions per category"
    chart.y_axis.title = "Users"
    chart.legend = None
    chart.add_data(Reference(summary, min_col=2, min_row=1, max_row=len(categories) + 1), titles_from_data=True)
    chart.set_categories(Reference(summary, min_col=1, min_row=2, max_row=len(categories) + 1))
    summary.add_chart(chart, "D2")

    workbook.save(OUTPUT_PATH)
    print(f"Saved Excel workbook to {OUTPUT_PATH.resolve()}")
    print("Open it in Excel and check both sheets: 'Quests' and 'By Category'.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
